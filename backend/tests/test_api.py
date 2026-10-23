"""API end to end with a fake vision model: upload -> process -> correct ->
export. Synthetic values only (iron rule 5)."""

import io
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.conftest import make_pdf
from vera.batch import Readers
from vera.config import get_settings
from vera.pipeline.vlm import PageReading, VlmResult

FIELDS = (
    "first_name last_name id_number central_numbers receipt_date sender "
    "commitment_numbers form_date sample_date"
).split()


def cand(value, label="", handwritten=False):
    return {
        "value_raw": value,
        "legible": True,
        "handwritten": handwritten,
        "label_nearby": label,
        "evidence_text": value,
        "region": "top-right",
    }


class FakeReader:
    """Same answer for every page: a clean form with one orange field."""

    def __init__(self, model):
        self.model = model

    def read_page(self, img, page_sha256, prompt):
        fields = {f: [] for f in FIELDS}
        fields.update(
            central_numbers=[cand("512340", "מרכזי")],
            commitment_numbers=[cand("88120", "התחייבות")],
            receipt_date=[cand("20/08/2026", "התקבל")],
            id_number=[cand("012345675", "ת.ז")],  # bad check digit -> review
            first_name=[cand("רות", "שם פרטי", True)],
            last_name=[cand("כהן", "שם משפחה", True)],
        )
        reading = PageReading.model_validate(
            {
                "page_kind": "form",
                "upright": True,
                "sender_guess": "שיבא",
                "page_marker": "",
                "fields": fields,
            }
        )
        return VlmResult(reading, self.model, "test", "k", True)


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{(tmp_path / 't.db').as_posix()}")
    monkeypatch.setenv("CPU_WORKERS", "1")
    get_settings.cache_clear()
    from vera.api.main import app

    with TestClient(app) as c:
        svc = app.state.batches
        svc.dpi, svc.orient = 60, False
        svc._readers = Readers(FakeReader("m1"), FakeReader("m2"), "prompt")
        yield c
    get_settings.cache_clear()


def wait_done(client, bid, timeout=60):
    t0 = time.time()
    while time.time() - t0 < timeout:
        b = client.get(f"/api/batches/{bid}").json()
        if b["status"] == "done":
            return b
        time.sleep(0.2)
    raise TimeoutError(b)


def upload(client, bid, tmp_path, pages=2):
    pdf = make_pdf(tmp_path / "in.pdf", pages=pages).read_bytes()
    return client.post(
        f"/api/batches/{bid}/files",
        files={"files": ("forms.pdf", pdf, "application/pdf")},
        headers={"X-User": "tester"},
    )


def test_upload_process_correct_export(client, tmp_path):
    bid = client.post("/api/batches", json={"name": "t", "ref_date": "2026-09-01"})
    bid = bid.json()["id"]
    r = upload(client, bid, tmp_path)
    assert r.status_code == 200 and r.json()["added"][0]["pages"] == 2

    b = wait_done(client, bid)
    assert b["counts"]["pages"] == {"done": 2}
    assert len(b["submissions"]) == 2
    sub = b["submissions"][0]
    f = sub["fields"]
    # The orange ID sent the page to escalation; the second model agrees on the
    # labelled central number, so it is green.
    assert f["central_numbers"] == {"values": ["512340"], "status": "verified"}
    assert f["id_number"]["status"] == "review"  # bad check digit
    # the official name from the closed list (config/senders.yaml)
    assert f["sender"]["values"] == ['המרכז הרפואי ע"ש חיים שיבא - תל השומר']

    # two different models agreeing on the labeled date -> green
    assert f["receipt_date"] == {"values": ["2026-08-20"], "status": "verified"}

    sid, pid = sub["id"], sub["page_ids"][0]
    img = client.get(f"/api/pages/{pid}/image?crop=0,0,100,100&scale=2")
    assert img.status_code == 200 and img.headers["content-type"] == "image/jpeg"

    queue = client.get(f"/api/batches/{bid}/review-queue").json()
    assert queue[0]["status"] == "review" and queue[0]["field"] == "id_number"

    # export is blocked while an orange field is unapproved
    blocked = client.post(f"/api/batches/{bid}/export")
    assert blocked.status_code == 409
    assert len(blocked.json()["detail"]["fields"]) == 2

    bad = client.patch(
        f"/api/submissions/{sid}/fields/id_number", json={"values": ["7"]}
    )
    assert bad.status_code == 422
    for s in b["submissions"]:
        ok = client.patch(
            f"/api/submissions/{s['id']}/fields/id_number",
            json={"values": ["12345674"]},
            headers={"X-User": "tester"},
        )
        assert ok.status_code == 200
    fr = ok.json()["fields"]["id_number"]
    assert fr["values"] == ["012345674"] and fr["status"] == "human"
    assert fr["history"][0]["user"] == "tester"

    assert client.get("/api/search?q=512340").json()[0]["field"] == "central_numbers"

    x = client.post(f"/api/batches/{bid}/export")
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    ws = wb["התחייבויות"]
    assert ws.sheet_view.rightToLeft
    assert ws.cell(row=2, column=4).value == "012345674"  # leading zero kept
    assert ws.cell(row=2, column=4).number_format == "@"
    assert wb["ביקורת"].max_row == 3  # header + 2 corrections

    # exported batch is locked
    again = client.patch(
        f"/api/submissions/{sid}/fields/first_name", json={"values": ["דנה"]}
    )
    assert again.status_code == 409


def test_non_document_upload_is_rejected(client):
    bid = client.post("/api/batches", json={"name": "t"}).json()["id"]
    r = client.post(
        f"/api/batches/{bid}/files",
        files={"files": ("x.pdf", b"PK\x03\x04 not a pdf", "application/pdf")},
    )
    assert r.status_code == 200
    assert r.json()["added"] == [] and r.json()["rejected"]
    assert r.json()["job_id"] is None


def test_unknown_batch_and_page_are_404(client):
    assert client.get("/api/batches/nope").status_code == 404
    assert client.get("/api/pages/nope/image").status_code == 404
    assert client.get("/api/submissions/nope").status_code == 404


def test_sidebar_count_is_the_open_batch_and_drops_on_approval(client, tmp_path):
    def waiting_on_screen(b):  # what the review screen counts: yellow + orange
        return sum(
            1
            for s in b["submissions"]
            for f in s["fields"].values()
            if f["status"] in ("likely", "review")
        )

    old = client.post("/api/batches", json={"name": "old", "ref_date": "2026-09-01"})
    upload(client, old.json()["id"], tmp_path)
    wait_done(client, old.json()["id"])
    bid = client.post("/api/batches", json={"name": "t", "ref_date": "2026-09-01"})
    bid = bid.json()["id"]
    upload(client, bid, tmp_path)
    b = wait_done(client, bid)

    o = client.get(f"/api/overview?batch_id={bid}").json()
    assert o["review_waiting_batch"] == waiting_on_screen(b) > 0
    assert o["review_waiting"] > o["review_waiting_batch"]  # both open batches
    assert client.get("/api/overview").json()["review_waiting_batch"] is None

    sub = b["submissions"][0]
    field = next(
        f for f, v in sub["fields"].items() if v["status"] in ("likely", "review")
    )
    client.post(f"/api/submissions/{sub['id']}/fields/{field}/accept")
    after = client.get(f"/api/overview?batch_id={bid}").json()
    assert after["review_waiting_batch"] == o["review_waiting_batch"] - 1


def test_accept_all_greens_values_and_leaves_empty_fields(client, tmp_path):
    bid = client.post("/api/batches", json={"name": "t", "ref_date": "2026-09-01"})
    bid = bid.json()["id"]
    upload(client, bid, tmp_path)
    before = wait_done(client, bid)["submissions"]
    has_value = sum(
        1
        for s in before
        for f in s["fields"].values()
        if f["values"] and f["status"] in ("likely", "review")
    )

    r = client.post(f"/api/batches/{bid}/accept-all", headers={"X-User": "tester"})
    assert r.status_code == 200
    assert r.json()["accepted"] == has_value

    for s in client.get(f"/api/batches/{bid}").json()["submissions"]:
        for f in s["fields"].values():
            if f["values"]:
                assert f["status"] in ("verified", "human")
            else:  # nothing accepted into an empty field
                assert f["status"] != "human"

    assert client.post("/api/batches/nope/accept-all").status_code == 404


def test_a_whole_page_image_does_not_shrink_the_cached_page(client, tmp_path):
    from PIL import Image

    bid = client.post("/api/batches", json={"name": "t", "ref_date": "2026-09-01"})
    bid = bid.json()["id"]
    upload(client, bid, tmp_path)
    pid = wait_done(client, bid)["submissions"][0]["page_ids"][0]
    page = client.get(f"/api/pages/{pid}").json()
    small = client.get(f"/api/pages/{pid}/image?max_side=100")
    assert max(Image.open(io.BytesIO(small.content)).size) == 100
    # a crop after that still comes from the full-size page
    crop = client.get(
        f"/api/pages/{pid}/image?crop=0,0,{page['width']},50&max_side=4000"
    )
    assert Image.open(io.BytesIO(crop.content)).size == (page["width"], 50)
