import pytest
from openpyxl import load_workbook

from vera.groundtruth import COLUMNS, load_ground_truth, write_template

PAGES = [("docA", 1), ("docA", 2), ("docB", 1)]
SPLITS = {"docA": "dev", "docB": "test"}


def test_template_has_one_row_per_page(tmp_path):
    out = tmp_path / "gt.xlsx"
    write_template(PAGES, out, SPLITS, images_dir=tmp_path / "pages")
    rows = load_ground_truth(out)
    assert [(r["doc"], r["page"], r["split"]) for r in rows] == [
        ("docA", "1", "dev"),
        ("docA", "2", "dev"),
        ("docB", "1", "test"),
    ]
    ws = load_workbook(out)["ground_truth"]
    assert [c.value for c in ws[1]] == COLUMNS
    assert ws.sheet_view.rightToLeft
    link = ws.cell(row=2, column=COLUMNS.index("image") + 1).hyperlink.target
    assert link == "pages/docA/p01.png"


def test_template_refuses_to_overwrite(tmp_path):
    out = tmp_path / "gt.xlsx"
    write_template(PAGES, out, SPLITS)
    with pytest.raises(FileExistsError):
        write_template(PAGES, out, SPLITS)
    write_template(PAGES, out, SPLITS, force=True)


def test_text_cells_keep_leading_zeros_and_letters(tmp_path):
    out = tmp_path / "gt.xlsx"
    write_template(PAGES, out, SPLITS)
    wb = load_workbook(out)
    ws = wb["ground_truth"]
    col = {name: i + 1 for i, name in enumerate(COLUMNS)}
    # Synthetic values shaped like the hard cases (not real sample values).
    ws.cell(row=2, column=col["commitment_numbers"]).value = "000123456789;X100001"
    ws.cell(row=2, column=col["id_number"]).value = "012345678"
    ws.cell(row=2, column=col["receipt_date"]).value = "01/02/2026"
    wb.save(out)

    row = load_ground_truth(out)[0]
    assert row["commitment_numbers"] == "000123456789;X100001"
    assert row["id_number"] == "012345678"
    assert row["receipt_date"] == "01/02/2026"
    assert ws.cell(row=2, column=col["id_number"]).number_format == "@"


def test_loads_csv_export(tmp_path):
    p = tmp_path / "gt.csv"
    p.write_text("doc,page,id_number\ndocA,1, 001234567 \n", encoding="utf-8-sig")
    assert load_ground_truth(p) == [
        {"doc": "docA", "page": "1", "id_number": "001234567"}
    ]
