"""Excel export (spec 6): one row per commitment form.

Sheets: "התחייבויות" (RTL; ID/central/commitment as text so leading zeros
survive; receipt date as a real Excel date; unverified cells coloured by
status with the explanation as a cell comment), "ביקורת" (every manual
change) and "סיכום" (counts by sender and status).
After writing, `verify_export` reads the file back and compares it to the
submissions (spec 6.3: the Excel must equal the screen).
"""

import io
from collections import Counter
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

from vera import __version__
from vera.models import Submission

SHEET = "התחייבויות"
# (header, field or None for computed columns)
COLUMNS = [
    ("מס' שורה", None),
    ("שם פרטי", "first_name"),
    ("שם משפחה", "last_name"),
    ("ת.ז.", "id_number"),
    ("מספר מרכזי", "central_numbers"),
    ("תאריך קבלה", "receipt_date"),
    ("גורם שולח", "sender"),
    ("מספר התחייבות", "commitment_numbers"),
    ("סטטוס שורה", None),
    ("קובץ ועמוד מקור", None),
]
TEXT_FIELDS = {"id_number", "central_numbers", "commitment_numbers"}
FILL = {
    "likely": PatternFill("solid", fgColor="FFF2CC"),  # yellow
    "review": PatternFill("solid", fgColor="F8CBAD"),  # orange
    "missing": PatternFill("solid", fgColor="D9D9D9"),  # grey
}
STATUS_HE = {
    "verified": "מאומת",
    "likely": "סביר",
    "review": "לבדיקה",
    "missing": "חסר",
    "human": "אושר ידנית",
}
SEP = "; "


def cell_text(sub: Submission, field: str) -> str:
    return SEP.join(sub.fields[field].values)


def blocking_fields(subs: list[Submission]) -> list[tuple[str, str]]:
    """Orange fields that nobody approved: they block the export (6.3)."""
    return [
        (s.id, f) for s in subs for f, fr in s.fields.items() if fr.status == "review"
    ]


def build_workbook(
    subs: list[Submission],
    corrections: list[dict],
    sources: dict[str, str],
    batch_name: str,
    date_format: str = "dd/mm/yyyy",
    forced: bool = False,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.sheet_view.rightToLeft = True
    ws.append([h for h, _ in COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")

    for n, sub in enumerate(subs, start=1):
        r = n + 1
        for col, (_, field) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=col)
            cell.alignment = Alignment(horizontal="right")
            if field is None:
                continue
            fr = sub.fields[field]
            if field == "receipt_date" and len(fr.values) == 1:
                try:
                    cell.value = date.fromisoformat(fr.values[0])
                    cell.number_format = date_format
                except ValueError:  # forced free text
                    cell.value = fr.values[0]
            else:
                cell.value = cell_text(sub, field) or None
                if field in TEXT_FIELDS:
                    cell.number_format = "@"
            if fr.status in FILL:
                cell.fill = FILL[fr.status]
                if fr.explanation_he:
                    cell.comment = Comment(fr.explanation_he, "VERA")
        ws.cell(row=r, column=1, value=n)
        ws.cell(row=r, column=9, value=STATUS_HE.get(sub.row_status, sub.row_status))
        ws.cell(row=r, column=10, value=sources.get(sub.id, ", ".join(sub.page_ids)))
    for col, width in zip("ABCDEFGHIJ", (8, 14, 14, 12, 18, 12, 24, 22, 12, 26)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "B2"

    audit = wb.create_sheet("ביקורת")
    audit.sheet_view.rightToLeft = True
    audit.append(["מי", "מתי", "שורה", "שדה", "ערך קודם", "ערך חדש", "פעולה", "מקור"])
    row_no = {s.id: i for i, s in enumerate(subs, start=1)}
    for c in corrections:
        audit.append(
            [
                c["user"],
                (
                    c["at"].strftime("%d/%m/%Y %H:%M")
                    if isinstance(c["at"], datetime)
                    else str(c["at"])
                ),
                row_no.get(c["submission_id"], c["submission_id"]),
                c["field"],
                SEP.join(c["old_values"]),
                SEP.join(c["new_values"]),
                "אישור" if c["action"] == "accept" else "תיקון",
                "ידני" if c["via"] == "ui" else "דרך העוזרת",
            ]
        )

    summary = wb.create_sheet("סיכום")
    summary.sheet_view.rightToLeft = True
    summary.append(["אצווה", batch_name])
    summary.append(["נוצר", datetime.now().strftime("%d/%m/%Y %H:%M")])
    summary.append(["גרסת מערכת", __version__])
    summary.append(["שורות", len(subs)])
    if forced:
        summary.append(["⚠️", "יוצא בכל זאת עם שדות שלא אושרו (מסומנים בכתום)"])
    summary.append([])
    summary.append(["סטטוס שורה", "כמות"])
    for st, k in Counter(s.row_status for s in subs).most_common():
        summary.append([STATUS_HE.get(st, st), k])
    summary.append([])
    summary.append(["גורם שולח", "כמות"])
    senders = Counter(cell_text(s, "sender") or "לא ידוע" for s in subs)
    for name, k in senders.most_common():
        summary.append([name, k])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def verify_export(data: bytes, subs: list[Submission]) -> list[str]:
    """Read the file back; every exported cell must equal the submission."""
    ws = load_workbook(io.BytesIO(data))[SHEET]
    problems = []
    for n, sub in enumerate(subs, start=1):
        for col, (_, field) in enumerate(COLUMNS, start=1):
            if field is None:
                continue
            got = ws.cell(row=n + 1, column=col).value
            if isinstance(got, datetime):
                got = got.date().isoformat()
            want = cell_text(sub, field)
            if (got or "") != want:
                problems.append(f"row {n} {field}: file={got!r} db={want!r}")
    if ws.max_row - 1 != len(subs):
        problems.append(f"rows: file={ws.max_row - 1} db={len(subs)}")
    return problems
