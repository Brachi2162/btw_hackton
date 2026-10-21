"""Ground-truth table (spec 17.1): template creation and loading.

Filled by hand in Excel, so it is an .xlsx with every cell text-formatted.
A CSV opened in Excel silently drops leading zeros ("000302651707") and
garbles Hebrew; see DECISIONS D-003.
"""

import csv
import os
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

# Spec 17.1 columns, plus `split` (dev/test) and `image` (link to the page).
COLUMNS = [
    "doc",
    "page",
    "split",
    "page_kind",
    "sender",
    "central_numbers",
    "receipt_date",
    "commitment_numbers",
    "first_name",
    "last_name",
    "id_number",
    "notes",
    "image",
]
PAGE_KINDS = ["form", "screenshot", "lab_report", "continuation", "unknown"]
TEXT = "@"

INSTRUCTIONS = [
    (
        "מה ממלאים",
        "את מה שכתוב בטופס בפועל, אחרי שמסתכלים על התמונה (קישור בעמודה image).",
    ),
    ("כמה ערכים", "מפרידים בנקודה-פסיק: 947468;947469"),
    ("שדה שלא קיים", "משאירים ריק. ריק = לא מופיע בטופס."),
    ("לא קריא", "משאירים ריק ומסבירים ב-notes (למשל: מרכזי, ספרה רביעית לא קריאה)."),
    ("תאריך קבלה", "בפורמט dd/mm/yyyy, למשל 25/08/2026"),
    ("ת.ז.", "9 ספרות, כולל אפסים מובילים. כתב יד גובר על הדפסה."),
    ("מספרים", "בדיוק כמו בטופס, כולל אפסים מובילים ואותיות (X132767)."),
    ("גורם שולח", "המוסד הספציפי ביותר (שניידר, לא כללית)."),
    ("page_kind", "form / screenshot / lab_report / continuation / unknown"),
    ("split", "dev = מכוונים עליו. test = רק למדידה, לא מסתכלים על כישלונות כדי לתקן."),
    ("אל תשנו", "את העמודות doc, page, split, image ואת סדר העמודות."),
]


def load_splits(path: Path) -> dict[str, str]:
    """doc stem -> 'dev' | 'test' (eval/splits.yaml)."""
    data = yaml.safe_load(Path(path).read_text(encoding="utf-8")) or {}
    return {str(k): str(v) for k, v in data.get("splits", {}).items()}


def write_template(
    pages: list[tuple[str, int]],
    out: Path,
    splits: dict[str, str],
    images_dir: Path | None = None,
    force: bool = False,
) -> None:
    """One empty row per (doc, page). Refuses to overwrite hand-filled work."""
    out = Path(out)
    if out.exists() and not force:
        raise FileExistsError(f"{out} exists; it may hold manual work (use force)")

    wb = Workbook()
    ws = wb.active
    ws.title = "ground_truth"
    ws.sheet_view.rightToLeft = True
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")

    for doc, page in pages:
        ws.append([doc, str(page), splits.get(doc, "")] + [""] * (len(COLUMNS) - 3))
        row = ws.max_row
        if images_dir is not None:
            target = Path(images_dir) / doc / f"p{page:02d}.png"
            link = Path(os.path.relpath(target, out.parent)).as_posix()
            img = ws.cell(row=row, column=COLUMNS.index("image") + 1)
            img.value, img.hyperlink, img.style = "פתח עמוד", link, "Hyperlink"

    for col_idx, name in enumerate(COLUMNS, start=1):
        letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[letter].width = 28 if name == "notes" else 16
        for (cell,) in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            cell.number_format = TEXT  # keep leading zeros, no auto-dates
            cell.alignment = Alignment(horizontal="right")

    last = max(ws.max_row, 2)
    kinds = DataValidation(type="list", formula1='"' + ",".join(PAGE_KINDS) + '"')
    kinds.add(f"D2:D{last}")
    ws.add_data_validation(kinds)
    ws.freeze_panes = "C2"

    help_ws = wb.create_sheet("הוראות")
    help_ws.sheet_view.rightToLeft = True
    help_ws.column_dimensions["A"].width = 16
    help_ws.column_dimensions["B"].width = 90
    for row in INSTRUCTIONS:
        help_ws.append(row)
        help_ws.cell(row=help_ws.max_row, column=1).font = Font(bold=True)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # someone typed a number into a non-text cell
    return str(value).strip()


def load_ground_truth(path: Path) -> list[dict[str, str]]:
    """Rows as {column: text}; reads the .xlsx template or an exported .csv."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            return [
                {k: (v or "").strip() for k, v in r.items()} for r in csv.DictReader(f)
            ]

    ws = load_workbook(path, read_only=True, data_only=True)["ground_truth"]
    rows = ws.iter_rows(values_only=True)
    header = [_cell_text(h) for h in next(rows)]
    out = []
    for values in rows:
        row = {h: _cell_text(v) for h, v in zip(header, values) if h}
        if row.get("doc"):
            out.append(row)
    return out
